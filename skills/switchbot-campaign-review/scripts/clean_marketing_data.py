#!/usr/bin/env python3
"""Normalize campaign tables, report data quality, and create an auditable workbook."""

from __future__ import annotations

import argparse
import json
import math
import os
import re
import unicodedata
from pathlib import Path
from typing import Any, Mapping

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from common import (
    CANONICAL_COLUMNS,
    CHANNEL_NORMALIZATION,
    LOGGER,
    NUMERIC_COLUMNS,
    RATE_COLUMNS,
    SourceTable,
    configure_logging,
    dataframe_to_markdown,
    inventory_dataframe,
    load_mapping,
    map_columns,
    normalize_token,
    parse_number,
    scan_input_directory,
)

ERROR_PATTERN = re.compile(r"#(?:REF!|VALUE!|DIV/0!|N/A|NAME\?|NUM!|NULL!)", re.IGNORECASE)
TOTAL_PATTERN = re.compile(r"^(?:grand\s*)?(?:total|subtotal|合計|总计|總計|小计|小計|汇总|集計)$", re.IGNORECASE)


def normalize_campaign_identity(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).casefold()
    text = re.sub(r"\b(?:19|20)\d{2}\b", " ", text)
    text = re.sub(
        r"(?:campaign\s*review|campaign|キャンペーン|活动复盘|活動復盤|大促复盘|投放复盘|复盘|復盤)",
        " ",
        text,
    )
    return re.sub(r"[^0-9a-z\u3040-\u30ff\u3400-\u9fff]+", "", text)


def apply_campaign_identity(
    cleaned: pd.DataFrame,
    issues: list[dict[str, Any]],
    expected_campaign_name: str | None,
) -> tuple[pd.DataFrame, dict[str, Any]]:
    source_names = sorted(
        {
            str(value).strip()
            for value in cleaned.get("campaign", pd.Series(dtype="object")).dropna()
            if str(value).strip() and str(value).strip().lower() != "nan"
        }
    )
    metadata: dict[str, Any] = {
        "expected_campaign_name": expected_campaign_name,
        "source_campaign_names": source_names,
        "status": "not_requested",
    }
    cleaned["report_campaign_name"] = expected_campaign_name or pd.NA
    cleaned["source_campaign_names"] = " | ".join(source_names) if source_names else pd.NA
    if not expected_campaign_name:
        cleaned["campaign_identity_status"] = "not_requested"
        return cleaned, metadata

    expected_token = normalize_campaign_identity(expected_campaign_name)
    source_tokens = [normalize_campaign_identity(name) for name in source_names]
    matched = any(
        token
        and expected_token
        and (
            token == expected_token
            or (min(len(token), len(expected_token)) >= 4 and (token in expected_token or expected_token in token))
        )
        for token in source_tokens
    )
    if matched:
        metadata["status"] = "verified"
    elif source_names:
        metadata["status"] = "conflict"
        issues.append(issue(
            "Report campaign identity conflict",
            "Critical",
            "Material impact",
            "ALL",
            "ALL",
            len(cleaned),
            (
                f"Requested report campaign={expected_campaign_name}; "
                f"source campaigns={source_names}."
            ),
            "Kept source campaign values and blocked period/channel comparisons.",
            "The data must not be presented as belonging to the requested campaign until confirmed.",
        ))
    else:
        metadata["status"] = "unverified"
        issues.append(issue(
            "Report campaign identity unverified",
            "Warning",
            "Needs confirmation",
            "ALL",
            "ALL",
            len(cleaned),
            f"Requested report campaign={expected_campaign_name}; no populated source campaign field was found.",
            "Used the requested name as a report label only.",
            "Campaign identity is not confirmed by the source data.",
        ))
    cleaned["campaign_identity_status"] = metadata["status"]
    return cleaned, metadata


def resolve_product_knowledge_dir(explicit: Path | None = None) -> Path:
    if explicit:
        return explicit.expanduser().resolve()
    codex_home = os.environ.get("CODEX_HOME")
    if codex_home:
        return (Path(codex_home).expanduser() / "skills" / "product-knowledge").resolve()
    return (Path.home() / ".codex" / "skills" / "product-knowledge").resolve()


def normalize_product_alias(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).strip().casefold()
    return re.sub(r"[\s\u3000・_\-]+", "", text)


def apply_product_knowledge(
    cleaned: pd.DataFrame,
    issues: list[dict[str, Any]],
    product_knowledge_dir: Path | None,
) -> tuple[pd.DataFrame, dict[str, Any]]:
    """Resolve product names exactly through product_index.json; never fuzzy match."""

    raw = cleaned["product"].copy()
    nonblank = raw.notna() & raw.astype(str).str.strip().ne("")
    cleaned["product_raw"] = raw.where(nonblank, pd.NA)
    cleaned["product_id"] = pd.NA
    knowledge_dir = resolve_product_knowledge_dir(product_knowledge_dir)
    index_path = knowledge_dir / "outputs" / "product_index.json"
    metadata: dict[str, Any] = {
        "path": str(knowledge_dir),
        "index_path": str(index_path),
        "index_generated_at": None,
        "validation_status": "not_executed",
        "index_validation_status": "not_executed",
        "resolution_status": "not_executed",
        "product_dimension_present": bool(nonblank.any()),
        "recognition_rate": None,
        "recognized_rows": 0,
        "product_rows": int(nonblank.sum()),
        "unrecognized_products": [],
        "ambiguous_products": [],
        "product_analysis_allowed": True,
    }
    if not nonblank.any():
        metadata["validation_status"] = "not_applicable"
        metadata["resolution_status"] = "not_applicable"
        issues.append(issue(
            "Product knowledge check not applicable", "Info", "No material impact",
            "ALL", "ALL", 0, "No populated product dimension was found.",
            "Skipped product entity resolution.", "Pure channel analysis may continue.",
        ))
        return cleaned, metadata

    if not index_path.exists():
        metadata["validation_status"] = "critical"
        metadata["index_validation_status"] = "missing"
        metadata["product_analysis_allowed"] = False
        issues.append(issue(
            "Product knowledge index missing", "Critical", "Material impact",
            "ALL", "ALL", int(nonblank.sum()), f"Index not found: {index_path}",
            "Product-level analysis blocked.", "Campaign output cannot safely aggregate by product.",
        ))
        return cleaned, metadata

    try:
        index = json.loads(index_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        metadata["validation_status"] = "critical"
        metadata["index_validation_status"] = "unreadable"
        metadata["product_analysis_allowed"] = False
        issues.append(issue(
            "Product knowledge index unreadable", "Critical", "Material impact",
            "ALL", "ALL", int(nonblank.sum()), str(exc),
            "Product-level analysis blocked.", "Campaign output cannot safely aggregate by product.",
        ))
        return cleaned, metadata

    metadata["index_generated_at"] = index.get("generated_at")
    metadata["index_validation_status"] = index.get("validation_status", "unknown")
    metadata["validation_status"] = metadata["index_validation_status"]
    if metadata["index_validation_status"] != "pass":
        metadata["product_analysis_allowed"] = False
        issues.append(issue(
            "Product knowledge validation failed", "Critical", "Material impact",
            "ALL", "ALL", int(nonblank.sum()),
            f"Index validation_status={metadata['index_validation_status']}.",
            "Product-level analysis blocked.", "Resolve Product Knowledge Critical findings and rebuild the index.",
        ))
        return cleaned, metadata
    metadata["resolution_status"] = "pass"

    lookup: dict[str, set[str]] = {}
    for alias, targets in index.get("alias_lookup", {}).items():
        lookup.setdefault(normalize_product_alias(alias), set()).update(targets)
    for product in index.get("products", []):
        lookup.setdefault(normalize_product_alias(product.get("product_id")), set()).add(product.get("product_id"))
    ambiguous = {
        normalize_product_alias(item.get("alias"))
        for item in index.get("excluded_aliases", [])
        if item.get("usage_status") == "ambiguous"
    }

    resolved = {}
    unknown_names = set()
    ambiguous_names = set()
    for row_index in cleaned.index[nonblank]:
        product_raw = str(cleaned.at[row_index, "product_raw"]).strip()
        token = normalize_product_alias(product_raw)
        targets = lookup.get(token, set())
        if token in ambiguous or len(targets) > 1:
            ambiguous_names.add(product_raw)
            continue
        if len(targets) == 1:
            resolved[row_index] = next(iter(targets))
        else:
            unknown_names.add(product_raw)
    for row_index, product_id in resolved.items():
        cleaned.at[row_index, "product_id"] = product_id
        cleaned.at[row_index, "product"] = product_id

    metadata["recognized_rows"] = len(resolved)
    metadata["recognition_rate"] = len(resolved) / int(nonblank.sum()) if nonblank.any() else None
    metadata["unrecognized_products"] = sorted(unknown_names)
    metadata["ambiguous_products"] = sorted(ambiguous_names)
    if unknown_names:
        issues.append(issue(
            "Unknown product alias", "Critical", "Needs mapping",
            "ALL", "ALL", int(cleaned["product_raw"].isin(unknown_names).sum()),
            f"Unrecognized product values: {', '.join(sorted(unknown_names))}.",
            "Preserved product_raw; left product_id blank; product-level analysis blocked.",
            "Add an exact, sourced alias to Product Knowledge and rebuild its index.",
        ))
    if ambiguous_names:
        issues.append(issue(
            "Ambiguous product alias", "Critical", "Needs mapping",
            "ALL", "ALL", int(cleaned["product_raw"].isin(ambiguous_names).sum()),
            f"Ambiguous product values: {', '.join(sorted(ambiguous_names))}.",
            "Preserved product_raw; left product_id blank; product-level analysis blocked.",
            "Resolve the alias one-to-one in Product Knowledge.",
        ))
    if unknown_names or ambiguous_names:
        metadata["resolution_status"] = "critical"
        metadata["product_analysis_allowed"] = False
        cleaned.loc[nonblank, "product"] = cleaned.loc[nonblank, "product_raw"]
    return cleaned, metadata


def issue(
    issue_type: str,
    severity: str,
    repair_class: str,
    source_file: str,
    source_table: str,
    affected_rows: int,
    detail: str,
    action: str,
    impact: str,
) -> dict[str, Any]:
    return {
        "Issue Type": issue_type,
        "Severity": severity,
        "Repair Class": repair_class,
        "Source File": source_file,
        "Source Table": source_table,
        "Affected Rows": affected_rows,
        "Detail": detail,
        "Action Taken": action,
        "Decision Impact": impact,
    }


def normalize_percent_series(
    series: pd.Series,
    column: str,
    source_file: str,
    source_table: str,
    issues: list[dict[str, Any]],
) -> pd.Series:
    explicit_mask = (
        series.astype(str)
        .map(lambda value: unicodedata.normalize("NFKC", value))
        .str.strip()
        .str.endswith("%")
    )
    parsed = series.map(parse_number)
    parsed.loc[explicit_mask] = parsed.loc[explicit_mask] / 100

    implicit = parsed.loc[~explicit_mask & parsed.notna()]
    above_one = implicit.gt(1)
    if not implicit.empty and above_one.all() and implicit.le(100).all():
        parsed.loc[implicit.index] = implicit / 100
        issues.append(
            issue(
                "Percentage scale normalized",
                "Info",
                "Auto-fixed",
                source_file,
                source_table,
                len(implicit),
                f"{column} contained 0-100 values without percent signs.",
                "Divided values by 100.",
                "No material impact after normalization.",
            )
        )
    elif above_one.any():
        issues.append(
            issue(
                "Mixed percentage scale",
                "Warning",
                "Needs confirmation",
                source_file,
                source_table,
                int(above_one.sum()),
                f"{column} mixes values above and below 1.",
                "Preserved ambiguous numeric scale; explicit percentages were normalized.",
                "Rate comparisons may be unreliable; recalculated ratios take precedence.",
            )
        )
    return parsed


def normalize_channel(value: Any) -> Any:
    if pd.isna(value):
        return value
    text = str(value).strip()
    return CHANNEL_NORMALIZATION.get(normalize_token(text), text)


def identify_total_rows(frame: pd.DataFrame) -> pd.Series:
    mask = pd.Series(False, index=frame.index)
    for column in ("channel", "campaign", "product", "subchannel"):
        if column in frame.columns:
            mask |= frame[column].astype(str).str.strip().str.match(TOTAL_PATTERN, na=False)
    return mask


def clean_tables(
    tables: list[SourceTable],
    custom_mapping: Mapping[str, str] | None = None,
    default_currency: str | None = None,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Return normalized rows, issues, inventory, unmapped fields, and excluded rows."""

    issues: list[dict[str, Any]] = []
    unmapped_records: list[dict[str, Any]] = []
    excluded_frames: list[pd.DataFrame] = []
    cleaned_frames: list[pd.DataFrame] = []
    inventory = inventory_dataframe(tables)

    for table in tables:
        if table.error:
            issues.append(
                issue(
                    "Source read error",
                    "Critical",
                    "Material impact",
                    table.source_file,
                    table.source_table,
                    0,
                    table.error,
                    "Source excluded from analysis.",
                    "Analysis is incomplete until the file can be read.",
                )
            )
            continue
        if table.data.empty:
            issues.append(
                issue(
                    "Empty table",
                    "Info",
                    "No material impact",
                    table.source_file,
                    table.source_table,
                    0,
                    "The table contains no data rows.",
                    "Skipped.",
                    "No impact unless data was expected.",
                )
            )
            continue

        normalized, mapped, unmapped = map_columns(table.data, custom_mapping)
        for field in unmapped:
            unmapped_records.append(
                {
                    "Source File": table.source_file,
                    "Source Table": table.source_table,
                    "Source Field": field,
                    "Status": "Unmapped",
                    "Action": "Retained as extra__ field; confirm whether it should map to a canonical field.",
                }
            )
        if unmapped:
            issues.append(
                issue(
                    "Unmapped fields",
                    "Warning",
                    "Needs confirmation",
                    table.source_file,
                    table.source_table,
                    len(table.data),
                    ", ".join(unmapped),
                    "Fields retained with extra__ prefix.",
                    "Potential metrics may be omitted from standardized analysis.",
                )
            )

        normalized["_source_row"] = normalized.index.to_series().astype(int) + 2
        normalized["_source_file"] = table.source_file
        normalized["_source_table"] = table.source_table

        source_string_view = normalized.astype(str)
        error_cells = source_string_view.apply(
            lambda column: column.str.contains(ERROR_PATTERN, na=False)
        )
        error_rows = error_cells.any(axis=1)
        if error_rows.any():
            issues.append(
                issue(
                    "Formula error strings",
                    "Critical",
                    "Material impact",
                    table.source_file,
                    table.source_table,
                    int(error_rows.sum()),
                    "Spreadsheet error tokens such as #REF! or #DIV/0! were found.",
                    "Rows retained; recalculated base-field metrics take precedence.",
                    "Affected source formulas are not trustworthy.",
                )
            )

        comparison_columns = [column for column in normalized.columns if not column.startswith("_source")]
        duplicate_mask = normalized[comparison_columns].astype(str).duplicated(keep="first")
        if duplicate_mask.any():
            excluded = normalized.loc[duplicate_mask].copy()
            excluded["_exclusion_reason"] = "Exact duplicate"
            excluded_frames.append(excluded)
            issues.append(
                issue(
                    "Exact duplicate rows",
                    "Warning",
                    "Auto-fixed",
                    table.source_file,
                    table.source_table,
                    int(duplicate_mask.sum()),
                    "Rows were identical across normalized analytical fields.",
                    "Kept the first row and excluded later copies.",
                    "Prevents duplicated totals.",
                )
            )
            normalized = normalized.loc[~duplicate_mask].copy()

        total_mask = identify_total_rows(normalized)
        if total_mask.any():
            excluded = normalized.loc[total_mask].copy()
            excluded["_exclusion_reason"] = "Summary or total row"
            excluded_frames.append(excluded)
            issues.append(
                issue(
                    "Summary rows mixed with detail",
                    "Warning",
                    "Auto-fixed",
                    table.source_file,
                    table.source_table,
                    int(total_mask.sum()),
                    "Rows labeled Total/Subtotal/合計/总计 were found.",
                    "Excluded from row-level calculations and retained in Excluded_Rows.",
                    "Prevents double counting.",
                )
            )
            normalized = normalized.loc[~total_mask].copy()

        for column in NUMERIC_COLUMNS:
            if column not in normalized.columns:
                continue
            if column in RATE_COLUMNS:
                normalized[column] = normalize_percent_series(
                    normalized[column], column, table.source_file, table.source_table, issues
                )
            else:
                original = normalized[column].copy()
                normalized[column] = original.map(parse_number)
                invalid = original.notna() & original.astype(str).str.strip().ne("") & normalized[column].isna()
                if invalid.any():
                    issues.append(
                        issue(
                            "Invalid numeric values",
                            "Warning",
                            "Needs confirmation",
                            table.source_file,
                            table.source_table,
                            int(invalid.sum()),
                            f"{column} contains non-numeric values.",
                            "Values retained as blank in the canonical numeric field.",
                            f"{column} totals may be understated.",
                        )
                    )

        if "date" in normalized.columns:
            original_date = normalized["date"].copy()
            normalized["date"] = pd.to_datetime(original_date, errors="coerce")
            invalid_date = original_date.notna() & original_date.astype(str).str.strip().ne("") & normalized["date"].isna()
            if invalid_date.any():
                issues.append(
                    issue(
                        "Invalid dates",
                        "Warning",
                        "Needs confirmation",
                        table.source_file,
                        table.source_table,
                        int(invalid_date.sum()),
                        "Some date values could not be parsed.",
                        "Unparseable canonical dates were left blank.",
                        "Daily and period analysis may omit these rows.",
                    )
                )

        if "period" not in normalized.columns and "date" in normalized.columns:
            normalized["period"] = normalized["date"].dt.year.astype("Int64").astype(str).replace("<NA>", pd.NA)
        elif "period" in normalized.columns:
            def normalize_period(value: Any) -> Any:
                if pd.isna(value):
                    return pd.NA
                if isinstance(value, (int, float)) and float(value).is_integer():
                    return str(int(value))
                text = str(value).strip()
                return pd.NA if text.lower() in {"", "nan", "<na>"} else text

            normalized["period"] = normalized["period"].map(normalize_period)

        if "channel" in normalized.columns:
            normalized["channel"] = normalized["channel"].map(normalize_channel)
        else:
            issues.append(
                issue(
                    "Missing channel field",
                    "Critical",
                    "Material impact",
                    table.source_file,
                    table.source_table,
                    len(normalized),
                    f"Mapped fields: {', '.join(sorted(mapped.values())) or 'None'}",
                    "Rows retained, but channel analysis is unavailable.",
                    "Channel evaluation cannot use these rows.",
                )
            )

        if "currency" not in normalized.columns and default_currency:
            normalized["currency"] = default_currency
        if "currency" in normalized.columns:
            normalized["currency"] = normalized["currency"].astype(str).str.upper().str.strip().replace({"NAN": pd.NA})

        key_columns = [column for column in ("date", "campaign", "channel", "product") if column in normalized.columns]
        if len(key_columns) >= 2:
            suspected = normalized.duplicated(subset=key_columns, keep=False)
            if suspected.any():
                issues.append(
                    issue(
                        "Possible business duplicates",
                        "Warning",
                        "Needs confirmation",
                        table.source_file,
                        table.source_table,
                        int(suspected.sum()),
                        f"Rows share business keys: {', '.join(key_columns)}.",
                        "Rows retained because their metric values may represent valid detail.",
                        "Totals could be overstated if the rows are duplicate exports.",
                    )
                )

        cleaned_frames.append(normalized)

    if not cleaned_frames:
        raise ValueError("No readable non-empty campaign tables were found")

    cleaned = pd.concat(cleaned_frames, ignore_index=True, sort=False)
    for column in CANONICAL_COLUMNS:
        if column not in cleaned.columns:
            cleaned[column] = pd.NA
    ordered = CANONICAL_COLUMNS + sorted(
        column for column in cleaned.columns if column not in CANONICAL_COLUMNS
    )
    cleaned = cleaned[ordered]

    for scope_field, label in (
        ("currency", "Multiple currencies"),
        ("tax_basis", "Mixed tax basis"),
        ("attribution_window", "Mixed attribution windows"),
        ("conversion_definition", "Mixed conversion definitions"),
        ("agency_fee_basis", "Mixed agency-fee basis"),
    ):
        distinct = cleaned[scope_field].dropna().astype(str).str.strip()
        distinct = sorted(value for value in distinct.unique() if value and value.lower() != "nan")
        if len(distinct) > 1:
            issues.append(
                issue(
                    label,
                    "Critical" if scope_field == "currency" else "Warning",
                    "Material impact",
                    "ALL",
                    "ALL",
                    len(cleaned),
                    f"{scope_field}: {', '.join(distinct)}",
                    "Kept values separate and flagged comparability.",
                    "Monetary totals or period comparisons may not be comparable.",
                )
            )

    missing_channel = cleaned["channel"].isna()
    if missing_channel.any():
        issues.append(
            issue(
                "Missing channel values",
                "Warning",
                "Needs confirmation",
                "ALL",
                "ALL",
                int(missing_channel.sum()),
                "Canonical channel is blank.",
                "Rows retained under Unassigned in grouped analysis.",
                "Channel contribution is incomplete.",
            )
        )

    issues_frame = pd.DataFrame(issues)
    unmapped_frame = pd.DataFrame(unmapped_records)
    excluded_frame = pd.concat(excluded_frames, ignore_index=True, sort=False) if excluded_frames else pd.DataFrame()
    return cleaned, issues_frame, inventory, unmapped_frame, excluded_frame


def write_quality_report(
    path: Path,
    issues: pd.DataFrame,
    inventory: pd.DataFrame,
    language: str = "zh-CN",
) -> None:
    severity_counts = (
        issues["Severity"].value_counts().rename_axis("Severity").reset_index(name="Count")
        if not issues.empty
        else pd.DataFrame(columns=["Severity", "Count"])
    )
    if language == "zh-CN":
        content = [
            "# 数据质量报告",
            "",
            "## 管理层判断",
            "",
            (
                f"[Fact] 已扫描 {len(inventory)} 个数据表，记录 {len(issues)} 个数据质量问题。"
                if len(inventory)
                else "[Data Gap] 未盘点到任何源数据表。"
            ),
            "",
            "## 数据源清单",
            "",
            dataframe_to_markdown(inventory),
            "",
            "## 严重程度汇总",
            "",
            dataframe_to_markdown(severity_counts),
            "",
            "## 问题明细",
            "",
            dataframe_to_markdown(issues, max_rows=200),
            "",
            "## 修复边界",
            "",
            "- 原始文件未被修改。",
            "- 仅自动排除完全重复记录和明确标记的汇总行。",
            "- 歧义映射、百分比口径、币种、税口径和归因窗口继续作为显式限制保留。",
        ]
    else:
        content = [
            "# Data Quality Report",
            "",
            "## Executive assessment",
            "",
            (
                f"[Fact] Scanned {len(inventory)} table(s); recorded {len(issues)} quality issue(s)."
                if len(inventory)
                else "[Data Gap] No source tables were inventoried."
            ),
            "",
            "## Source inventory",
            "",
            dataframe_to_markdown(inventory),
            "",
            "## Severity summary",
            "",
            dataframe_to_markdown(severity_counts),
            "",
            "## Issue details",
            "",
            dataframe_to_markdown(issues, max_rows=200),
            "",
            "## Repair boundary",
            "",
            "- Original files were not modified.",
            "- Only exact duplicates and clearly labeled total rows were excluded automatically.",
            "- Ambiguous mappings, percentage scales, currencies, tax bases, and attribution windows remain explicit limitations.",
        ]
    path.write_text("\n".join(content), encoding="utf-8")


def style_workbook(path: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="17223B")
    header_font = Font(color="FFFFFF", bold=True)
    light_fill = PatternFill("solid", fgColor="EAF0F8")
    thin = Side(style="thin", color="DCE3ED")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        if sheet.max_row and sheet.max_column:
            sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
        sheet.row_dimensions[1].height = 30
        for column_index in range(1, min(sheet.max_column, 80) + 1):
            values = [
                str(sheet.cell(row=row, column=column_index).value or "")
                for row in range(1, min(sheet.max_row, 200) + 1)
            ]
            width = min(max(max((len(value) for value in values), default=8) + 2, 10), 42)
            sheet.column_dimensions[get_column_letter(column_index)].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=False)
        if sheet.title in {
            "Source_Inventory",
            "Quality_Issues",
            "Unmapped_Fields",
            "Metric_Definitions",
        }:
            for row_index, row in enumerate(sheet.iter_rows(min_row=2), start=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                sheet.row_dimensions[row_index].height = 48
        if sheet.title == "Quality_Issues" and sheet.max_row > 1:
            sheet["A1"].fill = light_fill
            sheet["A1"].font = Font(bold=True, color="17223B")

    if "Normalized_Data" in workbook.sheetnames:
        sheet = workbook["Normalized_Data"]
        headers = {cell.value: cell.column for cell in sheet[1]}
        for field in ("ctr", "cvr", "vtr", "completion_rate"):
            if field in headers:
                for cell in sheet.iter_cols(min_col=headers[field], max_col=headers[field], min_row=2):
                    for item in cell:
                        item.number_format = "0.00%"
        for field in ("budget", "spend", "revenue", "cpc", "cpm", "cpa"):
            if field in headers:
                for cell in sheet.iter_cols(min_col=headers[field], max_col=headers[field], min_row=2):
                    for item in cell:
                        item.number_format = "#,##0.00"
        if "date" in headers:
            for cell in sheet.iter_cols(min_col=headers["date"], max_col=headers["date"], min_row=2):
                for item in cell:
                    item.number_format = "yyyy-mm-dd"
    workbook.save(path)


def write_cleaned_workbook(
    path: Path,
    cleaned: pd.DataFrame,
    inventory: pd.DataFrame,
    issues: pd.DataFrame,
    unmapped: pd.DataFrame,
    excluded: pd.DataFrame,
) -> None:
    metric_definitions = pd.DataFrame(
        [
            ["CTR", "Clicks / Impressions"],
            ["CPC", "Spend / Clicks"],
            ["CPM", "Spend / Impressions × 1,000"],
            ["CVR", "Conversions / Clicks"],
            ["CPA", "Spend / Conversions"],
            ["ROAS", "Revenue / Spend"],
            ["Budget Utilization", "Spend / Budget"],
            ["YoY", "(Current - Previous) / Previous"],
        ],
        columns=["Metric", "Definition"],
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        cleaned.to_excel(writer, sheet_name="Normalized_Data", index=False)
        inventory.to_excel(writer, sheet_name="Source_Inventory", index=False)
        issues.to_excel(writer, sheet_name="Quality_Issues", index=False)
        unmapped.to_excel(writer, sheet_name="Unmapped_Fields", index=False)
        excluded.to_excel(writer, sheet_name="Excluded_Rows", index=False)
        metric_definitions.to_excel(writer, sheet_name="Metric_Definitions", index=False)
    style_workbook(path)


def clean_marketing_data(
    input_dir: Path,
    output_dir: Path,
    mapping_file: Path | None = None,
    currency: str | None = None,
    product_knowledge_dir: Path | None = None,
    expected_campaign_name: str | None = None,
    language: str = "zh-CN",
) -> dict[str, Any]:
    tables = scan_input_directory(input_dir)
    mapping = load_mapping(mapping_file)
    cleaned, issues, inventory, unmapped, excluded = clean_tables(tables, mapping, currency)
    issue_records = issues.to_dict("records") if not issues.empty else []
    cleaned, campaign_identity = apply_campaign_identity(
        cleaned, issue_records, expected_campaign_name
    )
    cleaned, product_knowledge = apply_product_knowledge(
        cleaned, issue_records, product_knowledge_dir
    )
    issues = pd.DataFrame(issue_records)
    intermediate = output_dir / "_intermediate"
    intermediate.mkdir(parents=True, exist_ok=True)

    normalized_csv = intermediate / "normalized_data.csv"
    cleaned.to_csv(normalized_csv, index=False, encoding="utf-8-sig", date_format="%Y-%m-%d")
    inventory.to_csv(intermediate / "source_inventory.csv", index=False, encoding="utf-8-sig")
    issues.to_csv(intermediate / "quality_issues.csv", index=False, encoding="utf-8-sig")
    unmapped.to_csv(intermediate / "unmapped_fields.csv", index=False, encoding="utf-8-sig")
    excluded.to_csv(intermediate / "excluded_rows.csv", index=False, encoding="utf-8-sig")

    workbook_path = output_dir / "cleaned_data.xlsx"
    write_cleaned_workbook(workbook_path, cleaned, inventory, issues, unmapped, excluded)
    quality_report = output_dir / "data_quality_report.md"
    write_quality_report(quality_report, issues, inventory, language)
    LOGGER.info("Wrote %d normalized rows to %s", len(cleaned), workbook_path)
    return {
        "normalized_csv": normalized_csv,
        "workbook": workbook_path,
        "quality_report": quality_report,
        "campaign_identity": campaign_identity,
        "product_knowledge": product_knowledge,
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, default=Path("."))
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--mapping-file", type=Path)
    parser.add_argument("--currency", help="Default currency when the source has no currency field")
    parser.add_argument("--product-knowledge-dir", type=Path)
    parser.add_argument("--language", default="zh-CN")
    parser.add_argument("--verbose", action="store_true")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    configure_logging(args.verbose)
    outputs = clean_marketing_data(
        args.input_dir.resolve(),
        args.output_dir.resolve(),
        args.mapping_file.resolve() if args.mapping_file else None,
        args.currency,
        args.product_knowledge_dir,
        None,
        args.language,
    )
    for name, path in outputs.items():
        print(f"{name}: {json.dumps(path, ensure_ascii=False) if isinstance(path, dict) else path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
