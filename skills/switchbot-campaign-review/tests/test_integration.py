from __future__ import annotations

import json
import shutil
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

SKILL_DIR = Path(__file__).resolve().parents[1]
SCRIPTS = SKILL_DIR / "scripts"
sys.path.insert(0, str(SCRIPTS))

from build_report import build_report


def create_test_product_knowledge(tmp_path: Path) -> Path:
    skill_dir = tmp_path / "product-knowledge"
    outputs = skill_dir / "outputs"
    outputs.mkdir(parents=True)
    aliases = {
        "Lock Ultra": "lock_ultra",
        "Hub 3": "hub_3",
        "Robot Vacuum K10+": "k10_plus",
        "Lock Series": "lock_series",
        "Hub Series": "hub_series",
        "All Products": "all_products",
        "Multiple Products": "multiple_products",
    }
    (outputs / "product_index.json").write_text(
        json.dumps({
            "generated_at": "2026-07-30T00:00:00+09:00",
            "validation_status": "pass",
            "products": [{"product_id": product_id} for product_id in sorted(set(aliases.values()))],
            "alias_lookup": {alias: [product_id] for alias, product_id in aliases.items()},
            "excluded_aliases": [],
        }),
        encoding="utf-8",
    )
    return skill_dir


def test_full_pipeline_preserves_sources_and_supports_non_prime_campaign(
    tmp_path: Path,
) -> None:
    input_dir = tmp_path / "autumn-launch-data"
    input_dir.mkdir()
    source_names = ["sample_campaign_data.csv", "sample_previous_campaign_data.csv"]
    for name in source_names:
        shutil.copy2(SKILL_DIR / "examples" / name, input_dir / name)
    original_bytes = {
        name: (input_dir / name).read_bytes()
        for name in source_names
    }

    output_dir = tmp_path / "review-output"
    product_knowledge_dir = create_test_product_knowledge(tmp_path)
    result = build_report(
        input_dir=input_dir,
        campaign_name="Autumn Smart Home Launch",
        output_dir=output_dir,
        language="zh-CN",
        product_knowledge_dir=product_knowledge_dir,
    )
    assert result == output_dir.resolve()
    for name in (
        "executive_summary.md",
        "full_analysis.md",
        "data_quality_report.md",
        "cleaned_data.xlsx",
        "channel_evaluation.csv",
        "trend_analysis.csv",
        "action_plan.csv",
        "report.html",
    ):
        assert (output_dir / name).is_file()
        assert (output_dir / name).stat().st_size > 0

    assert {
        path.name for path in (output_dir / "charts").glob("*.png")
    } >= {
        "channel_spend_share.png",
        "channel_spend_vs_revenue_contribution.png",
        "tracked_roas_by_channel.png",
    }
    assert not (output_dir / "charts" / "spend_revenue_by_period.png").exists()

    for name in source_names:
        assert (input_dir / name).read_bytes() == original_bytes[name]

    workbook = load_workbook(output_dir / "cleaned_data.xlsx", read_only=True)
    assert workbook.sheetnames == [
        "Normalized_Data",
        "Source_Inventory",
        "Quality_Issues",
        "Unmapped_Fields",
        "Excluded_Rows",
        "Metric_Definitions",
    ]
    pd.read_csv(output_dir / "channel_evaluation.csv")
    pd.read_csv(output_dir / "trend_analysis.csv")
    actions = pd.read_csv(output_dir / "action_plan.csv")
    assert actions["Success Metric"].notna().all()

    report = (output_dir / "full_analysis.md").read_text(encoding="utf-8")
    assert "Autumn Smart Home Launch" in report
    assert "报告活动身份冲突" in report
    assert "产品级分析及产品汇总已省略" not in report
    for marker in (
        "[Fact]",
        "[Insight]",
        "[Hypothesis]",
        "[Recommendation]",
        "[Data Gap]",
        "[Risk]",
    ):
        assert marker in report

    summary = json.loads(
        (output_dir / "_intermediate" / "metrics_summary.json").read_text(
            encoding="utf-8"
        )
    )
    assert any(
        "Report campaign identity conflicts" in reason
        for reason in summary["comparability_reasons"]
    )
    assert any(
        "duration/date granularity differs" in reason
        for reason in summary["comparability_reasons"]
    )
    evaluation = pd.read_csv(output_dir / "channel_evaluation.csv")
    assert evaluation["YoY Change"].eq("Not Comparable").all()
    assert evaluation["Confidence"].eq("Low").all()
    trends = pd.read_csv(output_dir / "trend_analysis.csv")
    assert set(trends["Trend"]) == {"Insufficient comparable history"}
    html = (output_dir / "report.html").read_text(encoding="utf-8")
    assert "Autumn Smart Home Launch 活动复盘报告" in html
    assert "charts/" in html


def test_missing_product_knowledge_omits_only_product_rollups(
    tmp_path: Path,
) -> None:
    input_dir = tmp_path / "monthly-media-data"
    input_dir.mkdir()
    shutil.copy2(
        SKILL_DIR / "examples" / "sample_campaign_data.csv",
        input_dir / "sample_campaign_data.csv",
    )
    output_dir = tmp_path / "review-output"

    build_report(
        input_dir=input_dir,
        campaign_name="Monthly Media Review",
        output_dir=output_dir,
        product_knowledge_dir=tmp_path / "missing-product-knowledge",
    )

    manifest = json.loads(
        (output_dir / "_intermediate" / "run_manifest.json").read_text(
            encoding="utf-8"
        )
    )
    assert manifest["status"] == "partial"
    assert manifest["campaign_identity"]["status"] == "conflict"
    assert manifest["product_knowledge"]["product_analysis_allowed"] is False
    product_metrics = pd.read_csv(
        output_dir / "_intermediate" / "product_metrics.csv"
    )
    assert product_metrics.empty
    assert (output_dir / "report.html").is_file()
    assert (output_dir / "channel_evaluation.csv").is_file()
    quality = (output_dir / "data_quality_report.md").read_text(encoding="utf-8")
    assert "Product knowledge index missing" in quality
    full_report = (output_dir / "full_analysis.md").read_text(encoding="utf-8")
    assert "产品级分析及产品汇总已省略" in full_report
