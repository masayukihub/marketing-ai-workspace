#!/usr/bin/env python3
"""Read-only inspection of the Product Knowledge workbooks for Phase 7."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


SKILL_DIR = Path("/Users/lai/.codex/skills/product-knowledge")
REFERENCE_DIR = SKILL_DIR / "references"
PRODUCT_IDS = {"lock_ultra", "daily_station"}
ALIASES = {"lock ultra", "daily station", "ロックultra", "スマートデイリーステーション"}
WORKBOOKS = [
    "product_master.xlsx",
    "product_facts.xlsx",
    "pricing.xlsx",
    "compatibility.xlsx",
    "product_specs.xlsx",
    "product_positioning.xlsx",
    "product_entities.xlsx",
    "product_aliases.xlsx",
]


def normalize(value: object) -> str:
    return str(value or "").strip().casefold()


def relevant(row: dict[str, object]) -> bool:
    haystack = " | ".join(normalize(value) for value in row.values())
    return any(product_id in haystack for product_id in PRODUCT_IDS) or any(alias in haystack for alias in ALIASES)


def inspect_workbook(path: Path) -> dict[str, object]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    result: dict[str, object] = {"path": str(path), "sheets": []}
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        header_values = next(rows, ())
        headers = [str(value).strip() if value is not None else f"column_{index + 1}" for index, value in enumerate(header_values)]
        matches = []
        for values in rows:
            row = {headers[index]: value for index, value in enumerate(values) if index < len(headers) and value not in (None, "")}
            if row and relevant(row):
                matches.append(row)
        result["sheets"].append(
            {
                "name": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "headers": headers,
                "matching_rows": matches,
            }
        )
    workbook.close()
    return result


def main() -> None:
    index = json.loads((SKILL_DIR / "outputs/product_index.json").read_text(encoding="utf-8"))
    products = [row for row in index.get("products", []) if row.get("product_id") in PRODUCT_IDS]
    report = {
        "product_index_generated_at": index.get("generated_at"),
        "products": products,
        "workbooks": [inspect_workbook(REFERENCE_DIR / name) for name in WORKBOOKS],
    }
    print(json.dumps(report, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
