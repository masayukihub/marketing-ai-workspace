from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

SKILL_DIR = Path(__file__).resolve().parents[1]


def test_entity_registry_builds_backward_compatible_index() -> None:
    result = subprocess.run(
        [
            sys.executable,
            str(SKILL_DIR / "scripts" / "build_product_index.py"),
            "--skill-dir",
            str(SKILL_DIR),
            "--as-of",
            "2026-07-31",
        ],
        text=True,
        capture_output=True,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    payload = json.loads((SKILL_DIR / "outputs" / "product_index.json").read_text(encoding="utf-8"))
    assert payload["schema_version"] == "1.1"
    assert "alias_lookup" in payload
    assert set(payload["identifier_lookup"]) == {"asin", "sku", "jan", "model_number"}
    assert {entity["product_id"] for entity in payload["entities"]} >= {"lock_ultra", "hub_3", "ai_mindclip"}


def test_entity_registry_has_required_ids() -> None:
    workbook = load_workbook(SKILL_DIR / "references" / "product_entities.xlsx", read_only=True)
    sheet = workbook["Data"]
    headers = [cell.value for cell in sheet[1]]
    values = [
        dict(zip(headers, row))
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if any(value not in (None, "") for value in row)
    ]
    assert all(row["entity_id"] and row["product_id"] and row["entity_type"] for row in values)
