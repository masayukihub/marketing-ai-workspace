from __future__ import annotations

import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

SKILL_DIR = Path(__file__).resolve().parents[1]


def copy_skill(tmp_path: Path) -> Path:
    target = tmp_path / "product-knowledge"
    shutil.copytree(SKILL_DIR, target, ignore=shutil.ignore_patterns("__pycache__"))
    return target


def run_script(skill_dir: Path, name: str, *args: str) -> subprocess.CompletedProcess:
    return subprocess.run(
        [sys.executable, str(skill_dir / "scripts" / name), "--skill-dir", str(skill_dir), *args],
        text=True,
        capture_output=True,
    )


def append_row(path: Path, values: list) -> None:
    workbook = load_workbook(path)
    workbook["Data"].append(values)
    workbook.save(path)


def test_initial_knowledge_base_builds_index() -> None:
    result = run_script(SKILL_DIR, "build_product_index.py", "--as-of", "2026-07-30")
    assert result.returncode == 0, result.stdout + result.stderr
    assert (SKILL_DIR / "outputs" / "product_index.json").is_file()


def test_lock_ultra_cannot_own_face_recognition(tmp_path: Path) -> None:
    skill = copy_skill(tmp_path)
    workbook = load_workbook(skill / "references" / "product_facts.xlsx")
    sheet = workbook["Data"]
    headers = [cell.value for cell in sheet[1]]
    owner_col = headers.index("capability_owner_product_id") + 1
    fact_col = headers.index("fact_name") + 1
    for row in range(2, sheet.max_row + 1):
        if "face" in str(sheet.cell(row, fact_col).value).lower():
            sheet.cell(row, owner_col).value = "lock_ultra"
            break
    workbook.save(skill / "references" / "product_facts.xlsx")
    report = skill / "outputs" / "test-conflicts.md"
    result = run_script(skill, "check_conflicts.py", "--as-of", "2026-07-30", "--output", str(report))
    assert result.returncode == 1
    assert "CAPABILITY_OWNER_INVALID" in report.read_text(encoding="utf-8")


def test_firmware_required_fact_cannot_be_unqualified_approved(tmp_path: Path) -> None:
    skill = copy_skill(tmp_path)
    workbook = load_workbook(skill / "references" / "product_facts.xlsx")
    sheet = workbook["Data"]
    headers = [cell.value for cell in sheet[1]]
    sheet.cell(2, headers.index("release_status") + 1).value = "firmware_required"
    sheet.cell(2, headers.index("review_status") + 1).value = "approved"
    workbook.save(skill / "references" / "product_facts.xlsx")
    report = skill / "outputs" / "test-conflicts.md"
    result = run_script(skill, "check_conflicts.py", "--as-of", "2026-07-30", "--output", str(report))
    assert result.returncode == 1
    assert "APPROVED_SUPPORT_CONFLICT" in report.read_text(encoding="utf-8")


def test_performance_claim_without_conditions_is_blocked(tmp_path: Path) -> None:
    skill = copy_skill(tmp_path)
    workbook = load_workbook(skill / "references" / "product_facts.xlsx")
    sheet = workbook["Data"]
    headers = [cell.value for cell in sheet[1]]
    sheet.cell(2, headers.index("fact_type") + 1).value = "speed"
    sheet.cell(2, headers.index("conditions") + 1).value = ""
    sheet.cell(2, headers.index("review_status") + 1).value = "approved"
    sheet.cell(2, headers.index("release_status") + 1).value = "released"
    workbook.save(skill / "references" / "product_facts.xlsx")
    report = skill / "outputs" / "test-validation.md"
    result = run_script(skill, "validate_product_data.py", "--output", str(report))
    assert result.returncode == 1
    assert "FACT_CONDITIONS_MISSING" in report.read_text(encoding="utf-8")


def test_activity_and_regular_prices_cannot_overlap_silently(tmp_path: Path) -> None:
    skill = copy_skill(tmp_path)
    path = skill / "references" / "pricing.xlsx"
    common = [
        "hub_3", "Amazon", "JP", "Sale Price", "", None, "JPY", "", "",
        "Campaign", "2026-07-01", "2026-08-01", "yes", "SRC-004", "5836",
        "pending_verification", "2026-07-30", "",
    ]
    first = common.copy()
    first[5] = 10000
    second = common.copy()
    second[5] = 12000
    append_row(path, first)
    append_row(path, second)
    report = skill / "outputs" / "test-conflicts.md"
    result = run_script(skill, "check_conflicts.py", "--as-of", "2026-07-30", "--output", str(report))
    assert result.returncode == 1
    assert "PRICE_INTERVAL_CONFLICT" in report.read_text(encoding="utf-8")


def test_duplicate_product_invalid_source_and_alias_one_to_many_are_blocked(tmp_path: Path) -> None:
    skill = copy_skill(tmp_path)
    master = skill / "references" / "product_master.xlsx"
    workbook = load_workbook(master)
    sheet = workbook["Data"]
    duplicate = [cell.value for cell in sheet[2]]
    sheet.append(duplicate)
    sheet.cell(sheet.max_row, [cell.value for cell in sheet[1]].index("source_id") + 1).value = "SRC-999"
    workbook.save(master)
    aliases = skill / "references" / "product_aliases.xlsx"
    append_row(aliases, ["Hub 3", "lock_ultra", "short_name", "en", "", "", "current", "SRC-000", "2026-07-30", "test"])
    validation = skill / "outputs" / "test-validation.md"
    conflicts = skill / "outputs" / "test-conflicts.md"
    assert run_script(skill, "validate_product_data.py", "--output", str(validation)).returncode == 1
    assert "PRODUCT_ID_DUPLICATE" in validation.read_text(encoding="utf-8")
    assert "SOURCE_REFERENCE_INVALID" in validation.read_text(encoding="utf-8")
    assert run_script(skill, "check_conflicts.py", "--as-of", "2026-07-30", "--output", str(conflicts)).returncode == 1
    assert "ALIAS_ONE_TO_MANY" in conflicts.read_text(encoding="utf-8")


def test_copy_linter_blocks_lock_face_speed_and_battery_claims(tmp_path: Path) -> None:
    skill = copy_skill(tmp_path)
    assert run_script(skill, "build_product_index.py", "--as-of", "2026-07-30").returncode == 0
    report = skill / "outputs" / "claim-lint.md"
    result = run_script(
        skill,
        "check_copy_claims.py",
        "--text",
        "SwitchBot Lock Ultraなら、顔認証で0.3秒解錠。電池は12ヶ月持続。",
        "--output",
        str(report),
    )
    content = report.read_text(encoding="utf-8")
    assert result.returncode == 1
    assert "CAPABILITY_OWNER_COPY_ERROR" in content
    assert "NUMERIC_PERFORMANCE_REQUIRES_APPROVED_CLAIM" in content
    assert "PRODUCT_NOT_PUBLICATION_READY" in content


def test_copy_linter_blocks_chinese_owner_universal_and_no1_claims(tmp_path: Path) -> None:
    skill = copy_skill(tmp_path)
    assert run_script(skill, "build_product_index.py", "--as-of", "2026-07-30").returncode == 0
    cases = [
        (
            "Lock Ultra 支持 3D 人脸识别。",
            "CAPABILITY_OWNER_COPY_ERROR",
        ),
        (
            "Hub 3 支持所有 Matter 产品。",
            "UNIVERSAL_COMPATIBILITY_CLAIM",
        ),
        (
            "写一句“日本第一的智能锁”广告文案。",
            "UNSUPPORTED_SUPERLATIVE_CLAIM",
        ),
    ]
    for index, (copy, expected_code) in enumerate(cases, 1):
        report = skill / "outputs" / f"claim-lint-{index}.md"
        result = run_script(
            skill, "check_copy_claims.py",
            "--text", copy, "--output", str(report),
        )
        assert result.returncode == 1
        assert expected_code in report.read_text(encoding="utf-8")
