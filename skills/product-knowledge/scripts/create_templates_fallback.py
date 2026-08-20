#!/usr/bin/env python3
"""Create Product Knowledge workbooks when artifact-tool cannot load.

This fallback consumes the exact schema and seed rows defined in
create_templates.mjs, so there is only one template specification.
"""

from __future__ import annotations

import json
import subprocess
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


NODE = Path(
    "/Users/lai/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node"
)
HEADER_FILL = PatternFill("solid", fgColor="173F5F")
DICTIONARY_FILL = PatternFill("solid", fgColor="20639B")
ENUM_FILL = PatternFill("solid", fgColor="3CAEA3")
WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")
CRITICAL_FILL = PatternFill("solid", fgColor="F4CCCC")
OK_FILL = PatternFill("solid", fgColor="D9EAD3")
THIN = Side(style="thin", color="D9E2F3")
DEFAULT_FONT = "Hiragino Sans GB"


def load_spec(script_path: Path, skill_dir: Path) -> dict:
    command = [str(NODE), str(script_path), str(skill_dir), "--emit-json"]
    result = subprocess.run(command, check=True, capture_output=True, text=True)
    return json.loads(result.stdout)


def to_excel_value(field: str, value, date_fields: set[str]):
    if field not in date_fields or not value:
        return value
    if isinstance(value, (date, datetime)):
        return value
    return datetime.strptime(str(value), "%Y-%m-%d").date()


def configure_sheet(ws, headers: list[str], rows: list[list], enums: dict, date_fields: set[str], index: int):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 2 if len(headers) > 18 else 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"
    ws.page_margins.left = 0.2
    ws.page_margins.right = 0.2
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 2)}"
    ws.append(headers)
    for row in rows:
        ws.append([to_excel_value(field, value, date_fields) for field, value in zip(headers, row)])

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(name=DEFAULT_FONT, color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 34

    for row in ws.iter_rows(min_row=1, max_row=max(ws.max_row, 2), max_col=len(headers)):
        for cell in row:
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            if cell.row > 1:
                cell.font = Font(name=DEFAULT_FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    data_last_row = max(ws.max_row, 500)
    for column_index, field in enumerate(headers, 1):
        letter = get_column_letter(column_index)
        if field in date_fields:
            for cell in ws[letter][1:]:
                cell.number_format = "yyyy-mm-dd"
        if "price" in field or field == "discount_amount":
            for cell in ws[letter][1:]:
                cell.number_format = '#,##0'
        if field == "discount_rate":
            for cell in ws[letter][1:]:
                cell.number_format = "0.0%"
        if field in enums:
            values = ",".join(enums[field])
            validation = DataValidation(type="list", formula1=f'"{values}"', allow_blank=True)
            validation.error = f"Use one of: {values}"
            validation.errorTitle = "Invalid enum value"
            validation.prompt = f"Allowed: {values}"
            validation.promptTitle = field
            validation.showErrorMessage = True
            validation.showInputMessage = True
            ws.add_data_validation(validation)
            validation.add(f"{letter}2:{letter}{data_last_row}")

    for status_field in ("review_status", "status", "release_status", "usage_status"):
        if status_field not in headers:
            continue
        letter = get_column_letter(headers.index(status_field) + 1)
        applies = f"$A$2:${get_column_letter(len(headers))}${data_last_row}"
        for value in ("conflict", "rejected", "unsupported", "discontinued", "ambiguous"):
            ws.conditional_formatting.add(
                applies, FormulaRule(formula=[f'${letter}2="{value}"'], fill=CRITICAL_FILL)
            )
        for value in ("pending_verification", "draft", "unknown", "firmware_required", "expired"):
            ws.conditional_formatting.add(
                applies, FormulaRule(formula=[f'${letter}2="{value}"'], fill=WARNING_FILL)
            )
        for value in ("verified", "approved", "released", "active", "current"):
            ws.conditional_formatting.add(
                applies, FormulaRule(formula=[f'${letter}2="{value}"'], fill=OK_FILL)
            )

    if rows:
        ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        table = Table(displayName=f"DataTable{index}", ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False
        )
        ws.add_table(table)

    for column_index, field in enumerate(headers, 1):
        values = [len(str(field))]
        values.extend(
            len(str(ws.cell(row=row, column=column_index).value or ""))
            for row in range(2, min(ws.max_row, 80) + 1)
        )
        width = min(max(max(values) + 2, 11), 38)
        ws.column_dimensions[get_column_letter(column_index)].width = width


def add_dictionary(wb: Workbook, config: dict, enums: dict):
    ws = wb.create_sheet("Data Dictionary")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.append(["Field", "Required", "Allowed values", "Description"])
    required = set(config.get("required", []))
    descriptions = config.get("descriptions", {})
    for field in config["headers"]:
        ws.append([
            field,
            "Yes" if field in required else "No",
            " | ".join(enums.get(field, [])),
            descriptions.get(field, ""),
        ])
    for cell in ws[1]:
        cell.fill = DICTIONARY_FILL
        cell.font = Font(name=DEFAULT_FONT, color="FFFFFF", bold=True)
    for width, letter in zip((28, 12, 55, 80), ("A", "B", "C", "D")):
        ws.column_dimensions[letter].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row > 1:
                cell.font = Font(name=DEFAULT_FONT, size=10)
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def add_enums(wb: Workbook, enums: dict):
    ws = wb.create_sheet("Enums")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.append(["Field", "Allowed Value"])
    for field, values in enums.items():
        for value in values:
            ws.append([field, value])
    for cell in ws[1]:
        cell.fill = ENUM_FILL
        cell.font = Font(name=DEFAULT_FONT, color="FFFFFF", bold=True)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 36
    for row in ws.iter_rows():
        for cell in row:
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            if cell.row > 1:
                cell.font = Font(name=DEFAULT_FONT, size=10)


def main() -> int:
    skill_dir = Path(sys.argv[1] if len(sys.argv) > 1 else ".").resolve()
    script_path = skill_dir / "scripts" / "create_templates.mjs"
    spec = load_spec(script_path, skill_dir)
    reference_dir = skill_dir / "references"
    reference_dir.mkdir(parents=True, exist_ok=True)
    created = []
    for index, config in enumerate(spec["configs"], 1):
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        configure_sheet(
            ws, config["headers"], config.get("rows", []), spec["enums"],
            set(spec["isoDateColumns"]), index
        )
        add_dictionary(wb, config, spec["enums"])
        add_enums(wb, spec["enums"])
        target = reference_dir / config["file"]
        wb.save(target)
        created.append(target.name)
    print(json.dumps({"ok": True, "writer": "openpyxl-fallback", "workbooks": created}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
