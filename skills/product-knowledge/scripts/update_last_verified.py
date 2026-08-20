#!/usr/bin/env python3
"""Safely update verification dates for a source or product; dry-run by default."""

from __future__ import annotations

import argparse
import json
from datetime import date
from pathlib import Path

from openpyxl import load_workbook


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--skill-dir", required=True)
    selector = parser.add_mutually_exclusive_group(required=True)
    selector.add_argument("--source-id")
    selector.add_argument("--product-id")
    parser.add_argument("--source-revision-id", required=True)
    parser.add_argument("--date", default=date.today().isoformat())
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    skill_dir = Path(args.skill_dir).resolve()
    reference_dir = skill_dir / "references"
    changes = []
    for path in sorted(reference_dir.glob("*.xlsx")):
        workbook = load_workbook(path)
        if "Data" not in workbook.sheetnames:
            continue
        sheet = workbook["Data"]
        headers = [str(cell.value or "") for cell in sheet[1]]
        if "last_verified_date" not in headers:
            continue
        source_col = headers.index("source_id") + 1 if "source_id" in headers else None
        product_fields = [
            field for field in (
                "product_id", "our_product_id", "source_product_id",
                "subject_product_id", "capability_owner_product_id",
            ) if field in headers
        ]
        revision_col = headers.index("source_revision_id") + 1 if "source_revision_id" in headers else None
        date_col = headers.index("last_verified_date") + 1
        changed = 0
        for row in range(2, sheet.max_row + 1):
            source_match = args.source_id and source_col and str(sheet.cell(row, source_col).value or "") == args.source_id
            product_match = args.product_id and any(
                str(sheet.cell(row, headers.index(field) + 1).value or "") == args.product_id
                for field in product_fields
            )
            if not (source_match or product_match):
                continue
            before_date = sheet.cell(row, date_col).value
            before_revision = sheet.cell(row, revision_col).value if revision_col else None
            changes.append({
                "file": path.name, "row": row,
                "before_date": str(before_date or ""),
                "after_date": args.date,
                "before_revision": str(before_revision or ""),
                "after_revision": args.source_revision_id if revision_col else "not_applicable",
            })
            if args.apply:
                sheet.cell(row, date_col).value = date.fromisoformat(args.date)
                sheet.cell(row, date_col).number_format = "yyyy-mm-dd"
                if revision_col:
                    sheet.cell(row, revision_col).value = args.source_revision_id
            changed += 1
        if args.apply and changed:
            workbook.save(path)
    print(json.dumps({
        "ok": True,
        "mode": "apply" if args.apply else "dry-run",
        "selector": args.source_id or args.product_id,
        "source_revision_id": args.source_revision_id,
        "changes": changes,
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
