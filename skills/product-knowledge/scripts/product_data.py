#!/usr/bin/env python3
"""Shared readers and report helpers for Product Knowledge."""

from __future__ import annotations

import json
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

REVIEW_STATUS = {
    "draft", "pending_verification", "verified", "approved", "rejected",
    "expired", "conflict",
}
RELEASE_STATUS = {
    "released", "beta", "announced", "firmware_required", "region_limited",
    "unsupported", "unknown",
}
CLAIM_STATUS = {
    "Approved", "Conditional", "Internal Only", "Pending Verification",
    "Prohibited", "Expired",
}
SEVERITIES = ("Critical", "High", "Medium", "Low")


def value_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def load_rows(path: Path, sheet_name: str = "Data") -> tuple[list[str], list[dict[str, str]]]:
    workbook = load_workbook(path, data_only=False, read_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"{path.name}: missing sheet {sheet_name}")
    sheet = workbook[sheet_name]
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        raise ValueError(f"{path.name}: empty sheet")
    headers = [value_text(value) for value in values[0]]
    rows = []
    for row in values[1:]:
        record = {
            header: value_text(row[index]) if index < len(row) else ""
            for index, header in enumerate(headers)
        }
        if any(record.values()):
            rows.append(record)
    return headers, rows


def parse_source_registry(path: Path) -> dict[str, dict[str, str]]:
    text = path.read_text(encoding="utf-8")
    sources: dict[str, dict[str, str]] = {}
    blocks = re.split(r"(?m)^## ", text)
    for block in blocks[1:]:
        lines = block.splitlines()
        source_id = lines[0].strip()
        fields: dict[str, str] = {}
        for line in lines[1:]:
            match = re.match(r"^- ([^:]+):\s*(.*)$", line)
            if match:
                fields[match.group(1).strip()] = match.group(2).strip()
        sources[source_id] = fields
    return sources


def parse_profiles(products_dir: Path) -> dict[str, dict]:
    profiles: dict[str, dict] = {}
    for path in sorted(products_dir.glob("*.md")):
        text = path.read_text(encoding="utf-8")
        product_match = re.search(r"(?m)^- Product ID:\s*(\S+)", text)
        if not product_match:
            continue
        product_id = product_match.group(1)
        field = lambda name: (
            re.search(rf"(?m)^- {re.escape(name)}:[ \t]*(.*)$", text).group(1).strip()
            if re.search(rf"(?m)^- {re.escape(name)}:[ \t]*(.*)$", text) else ""
        )
        linked: dict[str, list[str]] = {}
        for label, value in re.findall(
            r"(?m)^- (Fact IDs|Claim IDs|Source IDs):[ \t]*(.*)$", text
        ):
            linked[label] = [
                item.strip() for item in value.split(";")
                if item.strip() and item.strip().lower() not in {"none", "pending verification"}
            ]
        profiles[product_id] = {
            "path": path,
            "completeness": field("Completeness"),
            "review_status": field("Review Status"),
            "market": field("Market"),
            "linked": linked,
        }
    return profiles


def parse_claims(path: Path) -> dict[str, dict[str, str]]:
    text = path.read_text(encoding="utf-8")
    claims: dict[str, dict[str, str]] = {}
    for block in re.split(r"(?m)^## ", text)[1:]:
        lines = block.splitlines()
        claim_match = re.search(r"\b(CLM-\d+)\b", lines[0])
        if not claim_match:
            continue
        claim_id = claim_match.group(1)
        fields = {}
        for line in lines[1:]:
            match = re.match(r"^- ([^:]+):\s*(.*)$", line)
            if match:
                fields[match.group(1).strip()] = match.group(2).strip()
        claims[claim_id] = fields
    return claims


def add_issue(issues: list[dict], severity: str, code: str, message: str, **context):
    issues.append({
        "severity": severity,
        "code": code,
        "message": message,
        "context": context,
    })


def counts(issues: list[dict]) -> dict[str, int]:
    counter = Counter(issue["severity"] for issue in issues)
    return {severity: counter.get(severity, 0) for severity in SEVERITIES}


def write_markdown_report(
    path: Path,
    title: str,
    issues: list[dict],
    metadata: dict | None = None,
):
    totals = counts(issues)
    lines = [
        f"# {title}",
        "",
        f"- Generated At: {datetime.now().astimezone().isoformat(timespec='seconds')}",
        f"- Overall Status: {'FAIL' if totals['Critical'] else 'PASS'}",
    ]
    if metadata:
        for key, value in metadata.items():
            lines.append(f"- {key}: {value}")
    lines.extend(["", "## Summary", ""])
    lines.extend(f"- {severity}: {totals[severity]}" for severity in SEVERITIES)
    lines.extend(["", "## Findings", ""])
    if not issues:
        lines.append("- No findings.")
    for issue in sorted(
        issues, key=lambda item: (SEVERITIES.index(item["severity"]), item["code"])
    ):
        context = "; ".join(
            f"{key}={value}" for key, value in issue.get("context", {}).items()
            if value not in ("", None, [])
        )
        suffix = f" ({context})" if context else ""
        lines.append(
            f"- [{issue['severity']}] `{issue['code']}` — {issue['message']}{suffix}"
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def dump_json(path: Path, payload: dict):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def normalize_alias(value: str) -> str:
    return re.sub(r"[\s\u3000・_\-]+", "", value).casefold()


def parse_iso(value: str) -> date | None:
    if not value:
        return None
    return date.fromisoformat(value[:10])
